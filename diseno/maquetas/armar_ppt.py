"""Arma la presentación del rediseño de ICEMM para stakeholders."""
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

BASE = pathlib.Path(__file__).parent
CAP = BASE / 'capturas'
LOGO = pathlib.Path(
    r'C:\Users\sleyton\BNV\Administración y Finanzas - Documentos'
    r'\Inteligencia de negocios\ICEMM\frontend\public\icemm-logo.png'
)
SALIDA = BASE / 'ICEMM - Rediseno de la app.pptx'

# Identidad EMM
NAVY   = RGBColor(0x10, 0x18, 0x20)
TEAL   = RGBColor(0x23, 0x30, 0x32)
MUTED  = RGBColor(0x80, 0x94, 0x94)
SAGE   = RGBColor(0x9D, 0xA3, 0x9B)
ACENTO = RGBColor(0xE0, 0x05, 0x44)
BORDE  = RGBColor(0xE8, 0xEB, 0xEF)
FONDO  = RGBColor(0xF7, 0xF8, 0xFA)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE  = RGBColor(0x2F, 0x7A, 0x55)

# Segoe UI existe en cualquier Windows. La marca usa Roboto, pero la app la
# carga de Google Fonts: no está garantizada en el equipo que presente, y un
# deck con sustitución automática se ve peor que uno con una neutral buena.
TIPO = 'Segoe UI'
TIPO_TIT = 'Segoe UI Semibold'

AN, AL = Inches(13.333), Inches(7.5)


def nueva():
    p = Presentation()
    p.slide_width, p.slide_height = AN, AL
    return p


def lamina(prs, fondo=FONDO):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.shapes.add_shape(1, 0, 0, AN, AL)
    f.fill.solid(); f.fill.fore_color.rgb = fondo
    f.line.fill.background(); f.shadow.inherit = False
    return s


def texto(s, x, y, an, al, txt, tam=14, color=TEAL, negrita=False,
          tipo=TIPO, espaciado=1.25, align=PP_ALIGN.LEFT, mayus=False):
    c = s.shapes.add_textbox(x, y, an, al)
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linea in enumerate(txt.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = espaciado
        r = p.add_run()
        r.text = linea.upper() if mayus else linea
        r.font.size = Pt(tam); r.font.bold = negrita
        r.font.color.rgb = color; r.font.name = tipo
    return c


def regla(s, x, y, an, al=Emu(22860), color=ACENTO):
    f = s.shapes.add_shape(1, x, y, an, al)
    f.fill.solid(); f.fill.fore_color.rgb = color
    f.line.fill.background(); f.shadow.inherit = False
    return f


def imagen_ajustada(s, ruta, cx, cy, an_max, al_max, borde=True):
    """Centra la imagen en la caja dada, conservando proporción."""
    w, h = Image.open(ruta).size
    esc = min(an_max / w, al_max / h)
    an, al = int(w * esc), int(h * esc)
    x, y = int(cx - an / 2), int(cy - al / 2)
    if borde:
        m = Emu(9525)
        f = s.shapes.add_shape(1, x - m, y - m, an + 2 * m, al + 2 * m)
        f.fill.solid(); f.fill.fore_color.rgb = BORDE
        f.line.fill.background(); f.shadow.inherit = False
    s.shapes.add_picture(str(ruta), x, y, an, al)
    return x, y, an, al


def pie(s, txt):
    texto(s, Inches(0.85), AL - Inches(0.62), Inches(11.6), Inches(0.3),
          txt, tam=9.5, color=SAGE)


def main():
    prs = nueva()

    # ─────────── 1. Portada ───────────
    s = lamina(prs, BLANCO)
    lw, lh = Image.open(LOGO).size
    an = Inches(2.5); al = int(an * lh / lw)
    s.shapes.add_picture(str(LOGO), Inches(0.85), Inches(0.8), an, al)
    regla(s, Inches(0.85), Inches(2.5), Inches(1.1), Emu(38100))
    texto(s, Inches(0.85), Inches(2.9), Inches(10), Inches(0.5),
          'Informe de Resultado de Obra', tam=15, color=MUTED)
    texto(s, Inches(0.85), Inches(3.35), Inches(11), Inches(1.4),
          'Rediseño de la aplicación', tam=46, color=NAVY, negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(4.95), Inches(9), Inches(0.9),
          'Propuesta de nuevas vistas para revisión.\nSeptiembre 2026',
          tam=13, color=TEAL, espaciado=1.5)
    pie(s, 'Las cifras que aparecen en las pantallas son de la demo: coherentes entre sí, '
           'no corresponden a una obra real.')

    # ─────────── 2. Qué cambia ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.75), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.95), Inches(11), Inches(0.7),
          'Qué cambia', tam=32, color=NAVY, negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(1.72), Inches(10.5), Inches(0.4),
          'Tres decisiones de fondo. Ninguna es cosmética: las tres salen de cómo se usa la app.',
          tam=13, color=MUTED)

    cols = [
        ('01', 'La navegación sube',
         'El menú lateral ocupaba 232 px fijos. La tabla de control tiene 174 cuentas y '
         'once columnas numéricas: es la pantalla donde esos píxeles más falta hacen.\n\n'
         'Con la barra superior el contenido pasa de 1.128 a 1.360 px útiles.'),
        ('02', 'El resultado, primero',
         'La pantalla abre con la cifra que decide: cuánto se proyecta ganar o perder al '
         'cierre, y contra qué presupuesto.\n\n'
         'El detalle por cuenta queda a un clic. Antes había que armar esa lectura a mano.'),
        ('03', 'Tema a elección',
         'Claro para sala con luz, proyector e impresión. Oscuro para trabajo prolongado '
         'frente a la pantalla.\n\n'
         'Lo elige cada usuario y queda guardado en su sesión. No es una decisión de la empresa.'),
    ]
    x0, anc, sep = Inches(0.85), Inches(3.5), Inches(0.42)
    for i, (n, tit, cuerpo) in enumerate(cols):
        x = x0 + i * (anc + sep)
        regla(s, x, Inches(2.45), anc, Emu(12700), BORDE)
        texto(s, x, Inches(2.72), anc, Inches(0.35), n, tam=11, color=ACENTO, negrita=True)
        texto(s, x, Inches(3.12), anc, Inches(0.5), tit, tam=19, color=NAVY,
              negrita=True, tipo=TIPO_TIT)
        texto(s, x, Inches(3.85), anc, Inches(2.6), cuerpo, tam=12, color=TEAL, espaciado=1.4)

    # ─────────── 3. Resumen de obra ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.6), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.78), Inches(8), Inches(0.5),
          'Resumen de obra', tam=26, color=NAVY, negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(1.26), Inches(11.6), Inches(0.35),
          'La pantalla de entrada: resultado proyectado, avance, semáforo por familia, '
          'curva de costo y las dos lecturas del cierre.', tam=12, color=MUTED)
    imagen_ajustada(s, CAP / 'resumen-claro.png', AN / 2, Inches(4.45),
                    Inches(11.5), Inches(5.0))
    pie(s, 'Tema claro.')

    # ─────────── 4. Costos ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.6), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.78), Inches(8), Inches(0.5),
          'Costos — tabla de control', tam=26, color=NAVY, negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(1.26), Inches(11.6), Inches(0.35),
          'Las once columnas se agrupan en Presupuesto · Ejecución · Resultado. '
          'La barra de la derecha muestra la desviación de cada familia sin leer un número.',
          tam=12, color=MUTED)
    imagen_ajustada(s, CAP / 'costos-claro.png', AN / 2, Inches(4.3),
                    Inches(11.5), Inches(4.6))
    pie(s, 'Las cuentas 900 (oficina central) quedan fuera del informe de obra.')

    # ─────────── 5. Los dos temas ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.6), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.78), Inches(9), Inches(0.5),
          'Claro y oscuro, a elección del usuario', tam=26, color=NAVY,
          negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(1.26), Inches(11.6), Inches(0.35),
          'Mismo sistema, misma tipografía, mismos datos. Cambia la superficie y el '
          'tono de las alertas, para que el rojo no vibre sobre fondo oscuro.',
          tam=12, color=MUTED)
    cy = Inches(4.3)
    imagen_ajustada(s, CAP / 'resumen-claro.png', Inches(3.5), cy, Inches(5.5), Inches(4.4))
    imagen_ajustada(s, CAP / 'resumen-oscuro.png', Inches(9.83), cy, Inches(5.5), Inches(4.4))
    texto(s, Inches(0.85), Inches(6.62), Inches(5.3), Inches(0.3), 'Claro',
          tam=11, color=TEAL, negrita=True, align=PP_ALIGN.CENTER)
    texto(s, Inches(7.18), Inches(6.62), Inches(5.3), Inches(0.3), 'Oscuro',
          tam=11, color=TEAL, negrita=True, align=PP_ALIGN.CENTER)

    # ─────────── 6. Construido vs propuesto ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.75), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.95), Inches(11), Inches(0.7),
          'Qué ya está construido y qué es propuesta', tam=30, color=NAVY,
          negrita=True, tipo=TIPO_TIT)
    texto(s, Inches(0.85), Inches(1.78), Inches(11), Inches(0.4),
          'Conviene separarlo, porque son dos cosas distintas y avanzan a distinta velocidad.',
          tam=13, color=MUTED)

    texto(s, Inches(0.85), Inches(2.65), Inches(5.4), Inches(0.35),
          'Construido y probado', tam=11, color=VERDE, negrita=True, mayus=True)
    regla(s, Inches(0.85), Inches(3.02), Inches(5.4), Emu(12700), VERDE)
    texto(s, Inches(0.85), Inches(3.25), Inches(5.4), Inches(3.2),
          'Exclusión de las cuentas 900 y de las contrapartidas de rollup al leer los '
          'archivos, con aviso de cuánto se descartó.\n\n'
          'Semáforo calculado por familia y no partida por partida.\n\n'
          'Directorio multi-obra con consolidado.\n\n'
          'Predictor de curvas de costo, con paridad verificada contra el prototipo '
          'dentro de 1 UF por mes.',
          tam=12, color=TEAL, espaciado=1.4)

    texto(s, Inches(7.1), Inches(2.65), Inches(5.4), Inches(0.35),
          'Propuesta, sin implementar', tam=11, color=ACENTO, negrita=True, mayus=True)
    regla(s, Inches(7.1), Inches(3.02), Inches(5.4), Emu(12700), ACENTO)
    texto(s, Inches(7.1), Inches(3.25), Inches(5.4), Inches(3.2),
          'Todo lo que se ve en estas láminas.\n\n'
          'Son maquetas: la aplicación hoy no se ve así. Lo que se decida acá pasa '
          'después a implementación.\n\n'
          'Faltan por dibujar tres pantallas: Familias, Directorio y la carga de archivos.',
          tam=12, color=TEAL, espaciado=1.4)

    # ─────────── 7. Qué sigue ───────────
    s = lamina(prs)
    regla(s, Inches(0.85), Inches(0.75), Inches(0.75))
    texto(s, Inches(0.85), Inches(0.95), Inches(11), Inches(0.7),
          'Qué sigue', tam=32, color=NAVY, negrita=True, tipo=TIPO_TIT)

    pasos = [
        ('Confirmar la dirección',
         'Barra superior y doble tema. Si hay acuerdo, el resto se dibuja sobre esa base.'),
        ('Completar las pantallas',
         'Familias, Directorio multi-obra y carga de archivos, en el mismo sistema.'),
        ('Implementar',
         'Sobre la rama de desarrollo, sin tocar lo que hoy está en producción.'),
        ('Lo que depende de EMM',
         'El costo por cuenta de las siete obras que faltan en el consolidado. Sin eso, '
         'la proyección no puede bajar de familia a cuenta individual.'),
    ]
    y = Inches(2.05)
    for i, (tit, cuerpo) in enumerate(pasos):
        texto(s, Inches(0.85), y, Inches(0.5), Inches(0.4), str(i + 1).zfill(2),
              tam=13, color=ACENTO, negrita=True)
        texto(s, Inches(1.65), y - Inches(0.03), Inches(4.2), Inches(0.4), tit,
              tam=16, color=NAVY, negrita=True, tipo=TIPO_TIT)
        texto(s, Inches(6.2), y, Inches(6.3), Inches(0.8), cuerpo, tam=12,
              color=TEAL, espaciado=1.35)
        if i < len(pasos) - 1:
            regla(s, Inches(0.85), y + Inches(0.92), Inches(11.63), Emu(12700), BORDE)
        y += Inches(1.22)

    prs.save(str(SALIDA))
    print('OK -', SALIDA.name, '-', len(prs.slides._sldIdLst), 'laminas')


if __name__ == '__main__':
    main()
